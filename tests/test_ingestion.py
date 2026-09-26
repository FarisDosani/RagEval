from uuid import UUID

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.ingestion import load_document
from app.schemas.document import ExtractedUnit


def test_txt_loading(tmp_path):
    path = tmp_path / "notes.TXT"
    text = "Hello, café!\nResearch notes.\n"
    path.write_text(text, encoding="utf-8")

    units = load_document(str(path))

    assert len(units) == 1
    unit = units[0]
    assert isinstance(unit, ExtractedUnit)
    assert unit.text == text
    assert unit.source == "notes.TXT"
    assert unit.page_number is None
    assert str(UUID(unit.document_id)) == unit.document_id


@pytest.mark.parametrize("text", ["", " \n\t"])
def test_empty_txt(tmp_path, text):
    path = tmp_path / "empty.txt"
    path.write_text(text, encoding="utf-8")
    assert load_document(path) == []


def test_unsupported_extension(tmp_path):
    with pytest.raises(ValueError, match="Unsupported document extension.*\\.csv"):
        load_document(tmp_path / "data.csv")


def test_missing_file(tmp_path):
    with pytest.raises(FileNotFoundError):
        load_document(tmp_path / "missing.txt")


def test_docx_paragraphs_and_metadata(tmp_path):
    path = tmp_path / "notes.docx"
    document = Document()
    for text in ["First paragraph", "", "  ", "Second paragraph"]:
        document.add_paragraph(text)
    document.save(path)

    units = load_document(path)

    assert [unit.text for unit in units] == ["First paragraph", "Second paragraph"]
    assert {unit.source for unit in units} == {path.name}
    assert {unit.page_number for unit in units} == {None}
    assert len({unit.document_id for unit in units}) == 1
    assert str(UUID(units[0].document_id)) == units[0].document_id


def test_empty_docx(tmp_path):
    path = tmp_path / "empty.docx"
    Document().save(path)
    assert load_document(path) == []


def test_pdf_pages_and_metadata(tmp_path):
    path = tmp_path / "pages.pdf"
    writer = PdfWriter()
    for text in ["First page", None, "Third page"]:
        page = writer.add_blank_page(width=200, height=200)
        if text is not None:
            font = DictionaryObject({
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            })
            page[NameObject("/Resources")] = DictionaryObject({
                NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})
            })
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode("ascii"))
            page[NameObject("/Contents")] = stream
    writer.write(path)

    units = load_document(path)

    assert [unit.text.strip() for unit in units] == ["First page", "Third page"]
    assert [unit.page_number for unit in units] == [1, 3]
    assert {unit.source for unit in units} == {path.name}
    assert len({unit.document_id for unit in units}) == 1
    assert str(UUID(units[0].document_id)) == units[0].document_id


def test_empty_pdf(tmp_path):
    path = tmp_path / "empty.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.write(path)
    assert load_document(path) == []
